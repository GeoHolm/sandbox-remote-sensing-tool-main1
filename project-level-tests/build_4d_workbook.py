#!/usr/bin/env python3
"""Fill the OpenET field data workbook for the 4D farm, one tab per field.

    python project-level-tests/build_4d_workbook.py

Starts from the blank template, keeps its layout, formatting and wording, and
replaces the two generic field tabs with four real ones -- the irrigated layout
for the two pivots, the rainfed layout for the two dryland fields.

PROVENANCE IS THE WHOLE POINT OF THIS SCRIPT
The form asks the grower what happened. Only grower records go in the shaded
cells: operation dates, implements and yields read out of the 4D export's
as-applied logs, ReadMe files, operations log, harvest report and irrigation
sheet. Satellite estimates never go in a shaded cell -- they sit in the Notes
column prefixed "SATELLITE:" so a reader can never mistake one for the other.
Where neither source has an answer the cell is left blank and the Notes column
says which source was checked.
"""
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding="utf-8")

HERE = Path(__file__).resolve().parent
PROJ = HERE / "4d-farm"
SRC = PROJ / "data to fill in" / "Farm XX - FTM_EarthRISE_OpenET_Field_Data_Collection.xlsx"
OUT = PROJ / "4D Farm - FTM_EarthRISE_OpenET_Field_Data_Collection.xlsx"

FARM = "4D Farm (DATA Farm), ABAC"
COUNTY, STATE = "Tift County", "Georgia"

# --------------------------------------------------------------- the data --
# acres from the boundary the pipeline measured; bbox for the location section.
FIELDS = {
    "North Pivot": dict(
        tab="North Pivot", layout="irrigated", key="north-pivot",
        acres=13.8, lat=31.4839, lon=-83.5479,
        bbox="31.48311 to 31.48472 N, -83.54963 to -83.54614 W",
        system="Center pivot", source="Not sure"),
    "South Pivot": dict(
        tab="South Pivot", layout="irrigated", key="south-pivot",
        acres=11.9, lat=31.4823, lon=-83.5479,
        bbox="31.48125 to 31.48305 N, -83.54950 to -83.54616 W",
        system="Center pivot", source="Not sure"),
    "Front Field": dict(
        tab="Front Field", layout="rainfed", key="front-field",
        acres=6.3, lat=31.4844, lon=-83.5449,
        bbox="31.48375 to 31.48497 N, -83.54602 to -83.54355 W",
        system="", source=""),
    "West Field": dict(
        tab="West Field", layout="rainfed", key="west-field",
        acres=16.6, lat=31.4821, lon=-83.5505,
        bbox="31.48072 to 31.48372 N, -83.55235 to -83.54836 W",
        system="", source=""),
}

B = "GROWER RECORD"
# crop year -> crop, yield, units, tillage, timing, extra passes, cover, notes
CROPS = {
 "north-pivot": {
  2022: dict(notes="No operation records in the 4D export for 2022. "
                   "SATELLITE: CDL says Cotton (100% of the boundary); "
                   "planting approx. 25 May, harvest approx. 14 Oct, low "
                   "confidence. Off-season before this crop not covered by the "
                   "satellite run."),
  2023: dict(notes="No operation records in the 4D export for 2023 on this "
                   "field. SATELLITE: CDL says Peanuts (77%); planting approx. "
                   "17 Jun, harvest approx. 31 Oct, medium confidence; winter "
                   "2022-23 read as green cover (peak NDVI 0.86)."),
  2024: dict(crop="Corn", yld=174.4, units="bu/ac", till="Strip till",
             timing="Multiple timings",
             extra="KMC ResiTill angled and straight passes 10 Nov 2023; "
                   "strip till 1st pass 19 Mar and 2nd pass 20 Mar 2024",
             cover="Cover crop — terminated before planting",
             notes="GROWER RECORD: crimson clover / cereal rye mix drilled "
                   "21 Nov 2023; burndown herbicide 11 Mar 2024; corn planted "
                   "20 Mar 2024; yield monitor mean 174.4 bu/ac over 3,782 "
                   "points; harvest 27 Aug 2024. "
                   "SATELLITE: planting approx. 8 Apr, harvest approx. 14 Sep, "
                   "low confidence — 19 days late on planting."),
  2025: dict(notes="No operation records in the 4D export for 2025. "
                   "SATELLITE: CDL says Cotton but only 48% of the boundary, "
                   "so treat the crop call as unreliable; planting approx. "
                   "31 May, harvest approx. 3 Oct; winter 2024-25 read as "
                   "likely cover crop (peak NDVI 0.73)."),
  2026: dict(notes="No operation records in the 4D export for 2026, and CDL "
                   "does not publish this year yet. SATELLITE: planting "
                   "approx. 8 Jun; no clear senescence yet so no harvest date; "
                   "winter 2025-26 read as likely cover crop (peak NDVI 0.85)."),
 },
 "south-pivot": {
  2022: dict(notes="No operation records in the 4D export for 2022. "
                   "SATELLITE: CDL says Peanuts but only 52% of the boundary, "
                   "so the crop call is unreliable; planting approx. 27 May, "
                   "harvest approx. 17 Oct, low confidence."),
  2023: dict(crop="Corn", yld=227.9, units="bu/ac", till="Reduced tillage",
             timing="Early spring (pre-plant)",
             extra="Single bedding pass 21 Mar 2023",
             cover="",
             notes="GROWER RECORD: dry fertilizer 20 Mar, bedding 21 Mar, "
                   "burndown spray 29 Mar, corn planted 30 Mar 2023 "
                   "(Dynagro D58VC65, 29,000 sd/ac); yield monitor mean "
                   "227.9 bu/ac over 4,096 points. No cover crop operation "
                   "recorded for the winter before this crop — the tillage "
                   "class is my reading of a single bedding pass, please "
                   "confirm. SATELLITE: planting approx. 10 Apr, harvest "
                   "approx. 2 Sep — 11 days late on planting."),
  2024: dict(crop="Cotton", yld=1.9, units="bales/ac", till="Strip till",
             timing="Spring (at planting)",
             extra="Strip till 2 May 2024",
             cover="Cover crop — terminated before planting",
             notes="GROWER RECORD: ryegrass drilled 6 Nov 2023, topdressed "
                   "13 Mar 2024, burndown 26 Apr, strip till 2 May, cotton "
                   "(Delta Pine DP2127) planted 13 May 2024 per the harvest "
                   "report; defoliant 7 Oct, picked 29 Oct. 21 bales of lint "
                   "over 11 acres = 1.9 bales/ac; yield monitor gave 1.8. "
                   "DISCREPANCY: the planting ReadMe is dated 20 Mar 2024, "
                   "which matches the North Pivot corn planting exactly and "
                   "cannot be right here — a strip till pass follows it on "
                   "2 May, and 20 Mar is far too early for Georgia cotton. "
                   "13 May used. Please confirm. "
                   "SATELLITE: planting approx. 3 Jun (window 19 May-18 Jun, "
                   "which does not contain 13 May), harvest approx. 27 Oct."),
  2025: dict(notes="No operation records in the 4D export for 2025. "
                   "SATELLITE: CDL says Peanuts (79%); planting approx. "
                   "29 May, harvest approx. 6 Oct, low confidence; winter "
                   "2024-25 read as likely cover crop (peak NDVI 0.90)."),
  2026: dict(notes="No operation records in the 4D export for 2026, and CDL "
                   "does not publish this year yet. SATELLITE: planting "
                   "approx. 12 Apr, harvest approx. 22 Sep; winter 2025-26 "
                   "read as likely cover crop (peak NDVI 0.86)."),
 },
 "front-field": {
  2022: dict(notes="No operation records in the 4D export for 2022. "
                   "SATELLITE: CDL says Cotton (97%); planting approx. 1 Jun, "
                   "harvest approx. 20 Oct, low confidence."),
  2023: dict(crop="Cotton", yld="", units="bales/ac", till="Strip till",
             timing="Early spring (pre-plant)",
             extra="Strip till 18 Apr 2023",
             cover="",
             notes="GROWER RECORD: lime 29 Mar, strip till 18 Apr, cotton "
                   "planted 15 May 2023, residual herbicide same day, "
                   "Pix growth regulator 6 Jul / 28 Jul / 11 Aug, defoliant "
                   "25 Aug. NO HARVEST FILE in the export, so no yield. No "
                   "cover crop operation recorded. "
                   "SATELLITE: planting approx. 25 May — 10 days late, and "
                   "15 May does fall inside the stated window."),
  2024: dict(crop="Peanuts", yld="", units="lb/ac", till="Strip till",
             timing="Spring (at planting)",
             extra="Strip till 2 May 2024; peanut inversion 30 Oct 2024 at "
                   "harvest",
             cover="Cover crop — terminated before planting",
             notes="GROWER RECORD: ryegrass drilled 7 Nov 2023, lime 29 Nov, "
                   "topdressed 13 Mar 2024, burndown 26 Apr, strip till 2 May, "
                   "inoculant and peanuts planted 6 Jun 2024 (Birdsong GA06G), "
                   "nine fungicide passes through the season, inverted "
                   "30 Oct. YIELD NOT AVAILABLE: the harvest file holds only "
                   "14 normalized points, not an absolute lb/ac. "
                   "SATELLITE: planting approx. 20 Jun — 14 days late; CDL "
                   "says Cotton at 73% purity, which disagrees with the "
                   "grower record of peanuts."),
  2025: dict(notes="No operation records in the 4D export for 2025. "
                   "SATELLITE: CDL says Soybeans (80%); planting approx. "
                   "6 Jun, harvest approx. 3 Oct; winter 2024-25 read as "
                   "likely cover crop (peak NDVI 0.77)."),
  2026: dict(notes="No operation records in the 4D export for 2026, and CDL "
                   "does not publish this year yet. SATELLITE: planting "
                   "approx. 12 Jun, harvest approx. 25 Sep; winter 2025-26 "
                   "read as green cover (peak NDVI 0.82)."),
 },
 "west-field": {
  2022: dict(notes="No operation records in the 4D export for 2022. "
                   "SATELLITE: CDL says Cotton (89%); planting approx. 1 Jun, "
                   "harvest approx. 19 Oct, low confidence."),
  2023: dict(crop="Cotton", yld="", units="bales/ac", till="",
             timing="", extra="", cover="",
             notes="GROWER RECORD: lime 26 Apr, cotton planted 16 May 2023, "
                   "residual herbicide 17 May, Pix 28 Jul and 11 Aug, "
                   "defoliant 25 Aug. NO TILLAGE OPERATION and NO HARVEST "
                   "FILE in the export for this field-year, so tillage and "
                   "yield are left blank rather than guessed. "
                   "SATELLITE: planting approx. 11 Jun — 26 days late, and "
                   "16 May falls outside the stated window, the worst season "
                   "date miss on this farm."),
  2024: dict(crop="Peanuts", yld=2752, units="lb/ac", till="Strip till",
             timing="Multiple timings",
             extra="Field cultivator 1 Dec 2023; KMC ResiTill 8 Dec 2023; "
                   "strip till 2 May 2024; peanut inversion 30 Oct 2024",
             cover="Cover crop — terminated before planting",
             notes="GROWER RECORD: lime 29 Nov 2023, field cultivator 1 Dec, "
                   "cereal rye drilled 7 Dec, ResiTill 8 Dec, burndown "
                   "26 Apr 2024, strip till 2 May, inoculant and peanuts "
                   "planted 6 Jun 2024 (Birdsong GA06G), nine fungicide "
                   "passes, inverted 30 Oct. Yield monitor mean 2,752 lb/ac "
                   "over 3,434 points at 17% moisture. "
                   "SATELLITE: planting approx. 25 Jun — 19 days late; CDL "
                   "says Cotton at 92% purity, which disagrees with the "
                   "grower record of peanuts."),
  2025: dict(notes="No operation records in the 4D export for 2025. "
                   "SATELLITE: CDL says Soybeans but only 57% of the "
                   "boundary, so treat the crop call as unreliable; planting "
                   "approx. 28 May, harvest approx. 27 Sep; winter 2024-25 "
                   "read as likely cover crop (peak NDVI 0.67)."),
  2026: dict(notes="No operation records in the 4D export for 2026, and CDL "
                   "does not publish this year yet. SATELLITE: planting "
                   "approx. 6 Jun, harvest approx. 25 Sep; winter 2025-26 "
                   "read as likely cover crop (peak NDVI 0.78)."),
 },
}

# Season-long inches applied, from "irrigation rates.xlsx" in the export.
IRRIG = {
 "north-pivot": {2024: (7.50, 10, "10 events at 0.75 in"),
                 2025: (0.00, 0, "No irrigation events recorded for 2025"),
                 2026: (6.25, 9, "2 events at 0.5 in and 7 at 0.75 in; "
                                 "39 days to maturity at termination")},
 "south-pivot": {2024: (6.75, 9, "9 events at 0.75 in"),
                 2025: (6.75, 9, "9 events at 0.75 in"),
                 2026: (8.75, 13, "4 events at 0.5 in and 9 at 0.75 in; "
                                  "crop matured, irrigation terminated")},
}
IRRIG_METHOD = ("Pivot event log kept by the farm, converted from event count "
                "x nominal depth. Valley irrigation scheduling with in-field "
                "soil moisture sensors.")

YEARS = [2022, 2023, 2024, 2025, 2026]


def fill(ws, coord, value, note_font=False):
    c = ws[coord]
    c.value = value
    if note_font:
        c.font = Font(name="Arial", size=9, italic=True, color="FF58595B")
    return c


def build():
    wb = openpyxl.load_workbook(SRC)
    src_irr, src_rain = wb["Field 1 (Irrigated)"], wb["Field 2 (Rainfed)"]

    for name, f in FIELDS.items():
        base = src_irr if f["layout"] == "irrigated" else src_rain
        ws = wb.copy_worksheet(base)
        ws.title = f["tab"]

        # --- 1. field information -----------------------------------------
        ws["A1"] = "%s  —  %s" % (f["tab"].upper(),
                                  "IRRIGATED" if f["layout"] == "irrigated"
                                  else "RAINFED")
        fill(ws, "B7", FARM)
        fill(ws, "B8", f["tab"])
        fill(ws, "B9", COUNTY)
        fill(ws, "B10", STATE)
        fill(ws, "B11", f["acres"])
        if f["layout"] == "irrigated":
            fill(ws, "B12", f["system"])
            fill(ws, "B13", f["source"])

        # --- 2. location ---------------------------------------------------
        loc_row = 17 if f["layout"] == "irrigated" else 15
        fill(ws, "A%d" % loc_row,
             "Boundary supplied as GeoJSON with this workbook "
             "(%s.geojson).\nCentroid  %.4f N, %.4f W\nBounding box  %s"
             % (f["key"], f["lat"], f["lon"], f["bbox"]))
        ws["A%d" % loc_row].alignment = Alignment(wrap_text=True,
                                                  vertical="top")

        # --- 3. crop, yield and tillage ------------------------------------
        first = 36 if f["layout"] == "irrigated" else 34
        for i, yr in enumerate(YEARS):
            r = first + i
            d = CROPS[f["key"]].get(yr, {})
            fill(ws, "A%d" % r, yr)
            fill(ws, "B%d" % r, d.get("crop", ""))
            fill(ws, "C%d" % r, d.get("yld", ""))
            fill(ws, "D%d" % r, d.get("units", "") if d.get("crop") else "")
            fill(ws, "E%d" % r, d.get("till", ""))
            fill(ws, "F%d" % r, d.get("timing", ""))
            fill(ws, "G%d" % r, d.get("extra", ""))
            fill(ws, "H%d" % r, d.get("cover", ""))
            fill(ws, "I%d" % r, d.get("notes", ""))
            for col in "GHI":
                ws["%s%d" % (col, r)].alignment = Alignment(wrap_text=True,
                                                            vertical="top")
            ws.row_dimensions[r].height = 92

        # --- 4. irrigation (pivots only) -----------------------------------
        if f["layout"] == "irrigated":
            # The irrigation block is four data rows, 48-51; row 52 is a merged
            # explanatory note, not a spare. Writing a fifth year into it
            # destroyed that note and silently dropped B-D into merged cells.
            # The export only has irrigation for 2024-2026 anyway, so the four
            # rows run 2023-2026.
            for i, yr in enumerate(YEARS[1:]):
                r = 48 + i
                fill(ws, "A%d" % r, yr)
                v = IRRIG[f["key"]].get(yr)
                if v:
                    inches, events, desc = v
                    fill(ws, "B%d" % r, inches)
                    fill(ws, "C%d" % r, events)
                    fill(ws, "D%d" % r, desc)
                    fill(ws, "I%d" % r, IRRIG_METHOD)
                else:
                    fill(ws, "D%d" % r,
                         "Not in the 4D export — the irrigation sheet covers "
                         "2024 to 2026 only.")
                ws["D%d" % r].alignment = Alignment(wrap_text=True,
                                                    vertical="top")
                ws["I%d" % r].alignment = Alignment(wrap_text=True,
                                                    vertical="top")
                ws.row_dimensions[r].height = 46

        # --- 5. anything else ----------------------------------------------
        last = 56 if f["layout"] == "irrigated" else 47
        extra = ("Prepared by Field to Market from the 4D Farm data export, "
                 "October 2026. Two sources, kept separate throughout:\n"
                 "  •  GROWER RECORD — as-applied logs, operation ReadMe "
                 "files, the 2023 operations log, the cotton harvest report "
                 "and the farm irrigation sheet. These are in the shaded "
                 "cells.\n"
                 "  •  SATELLITE — Sentinel-2 run of the Field to Market "
                 "remote sensing pipeline, 2022-2026. These appear only in "
                 "the Notes column, always prefixed, and are estimates rather "
                 "than observations.\n"
                 "A blank shaded cell means neither source answered it; the "
                 "Notes column says which was checked.\n"
                 "Please correct anything wrong — the point of sharing this "
                 "is that the farm knows and we are inferring.")
        if f["layout"] == "irrigated":
            extra += ("\nWater source is left as 'Not sure': the export shows "
                      "Valley pivot scheduling and soil moisture sensors but "
                      "never names the water source.")
        fill(ws, "A%d" % last, extra)
        ws["A%d" % last].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[last].height = 150

        # Dropdowns, rebuilt. openpyxl drops the template's originals on load
        # because they are stored in the unsupported extension form, so the
        # saved file would otherwise have no validation at all.
        for col, src, n in (("B", "$A$2:$A$17", 1), ("E", "$C$2:$C$8", 1),
                            ("F", "$D$2:$D$9", 1), ("H", "$G$2:$G$5", 1)):
            dv = DataValidation(type="list",
                                formula1="=Lists!%s" % src, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add("%s%d:%s%d" % (col, first, col, first + len(YEARS) - 1))
        if f["layout"] == "irrigated":
            for cell, src in (("B12", "$E$2:$E$10"), ("B13", "$F$2:$F$7")):
                dv = DataValidation(type="list",
                                    formula1="=Lists!%s" % src,
                                    allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(cell)

    del wb["Field 1 (Irrigated)"]
    del wb["Field 2 (Rainfed)"]
    order = ["About This Project"] + [f["tab"] for f in FIELDS.values()] + ["Lists"]
    wb._sheets = [wb[t] for t in order]

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("wrote %s" % OUT)
    print("tabs:", order)


if __name__ == "__main__":
    build()
